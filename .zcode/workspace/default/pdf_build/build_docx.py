# -*- coding: utf-8 -*-
"""Build a formatted DOCX from the Ronaldo vs Messi markdown research report."""
import re
import sys
import json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

MD_PATH = r"C:\Users\ilayl\.zcode\workspace\default\ronaldo_vs_messi_deep_research.md"
OUT_PATH = r"C:\Users\ilayl\.zcode\workspace\default\ronaldo_vs_messi_deep_research.docx"

ACCENT = RGBColor(0x1F, 0x3A, 0x5F)   # dark slate blue
GRAY = RGBColor(0x59, 0x59, 0x59)
LINK = RGBColor(0x0B, 0x5C, 0xAD)
BODY_FONT = "Calibri"

TOKEN_RE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|\[[^\]]+\]\([^)\s]+\))")


def add_hyperlink(paragraph, text, url, size=None, bold=False):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0B5CAD"); rPr.append(color)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rPr.append(u)
    if size:
        sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    if bold:
        b = OxmlElement("w:b"); rPr.append(b)
    run.append(rPr)
    t = OxmlElement("w:t"); t.set(qn("xml:space"), "preserve"); t.text = text
    run.append(t)
    hl.append(run)
    paragraph._p.append(hl)


def emit_inline(paragraph, text, size=None, base_bold=False, color=None):
    """Parse **bold**, *italic* and [text](url) into runs."""
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        m_link = re.fullmatch(r"\[([^\]]+)\]\(([^)\s]+)\)", tok)
        m_bold = tok.startswith("**") and tok.endswith("**")
        m_ital = tok.startswith("*") and tok.endswith("*") and not m_bold
        if m_link:
            add_hyperlink(paragraph, m_link.group(1), m_link.group(2), size=size, bold=base_bold)
        elif m_bold:
            r = paragraph.add_run(tok[2:-2])
            r.bold = True
            r.font.size = size; r.font.name = BODY_FONT
            if color: r.font.color.rgb = color
        elif m_ital:
            r = paragraph.add_run(tok[1:-1])
            r.italic = True
            r.font.size = size; r.font.name = BODY_FONT
            if color: r.font.color.rgb = color
        else:
            r = paragraph.add_run(tok)
            r.bold = base_bold
            r.font.size = size; r.font.name = BODY_FONT
            if color: r.font.color.rgb = color


def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def mark_header_row(row):
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)


def no_row_split(row):
    """Keep the row intact across page breaks (LibreOffice honours w:cantSplit)."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def table_full_width(table):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); tblPr.append(tblW)
    tblW.set(qn("w:w"), "5000"); tblW.set(qn("w:type"), "pct")


def add_footer_page_number(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.font.size = Pt(9); run.font.color.rgb = GRAY; run.font.name = BODY_FONT
    fld1 = OxmlElement("w:fldChar"); fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
    fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1); run._r.append(instr); run._r.append(fld2)


def build_title_page(doc, title, meta_lines):
    for _ in range(5):
        doc.add_paragraph()
    p_kicker = doc.add_paragraph()
    p_kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_kicker.add_run("ГЛУБОКОЕ СРАВНИТЕЛЬНОЕ ИССЛЕДОВАНИЕ")
    r.font.size = Pt(12); r.font.color.rgb = GRAY; r.bold = True; r.font.name = BODY_FONT

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(18)
    r = p_title.add_run(title)
    r.font.size = Pt(28); r.bold = True; r.font.color.rgb = ACCENT; r.font.name = BODY_FONT

    # accent line via paragraph border
    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = p_line._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "1F3A5F")
    pBdr.append(bottom); pPr.append(pBdr)

    doc.add_paragraph()
    for line in meta_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        emit_inline(p, line, size=Pt(11), color=GRAY)

    for _ in range(6):
        doc.add_paragraph()
    p_src = doc.add_paragraph()
    p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_src.add_run("Подготовлено на основе официальных протоколов FIFA, UEFA,\nнациональных лиг и клубов")
    r.font.size = Pt(10); r.font.color.rgb = GRAY; r.italic = True; r.font.name = BODY_FONT
    doc.add_page_break()


def style_setup(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), BODY_FONT)
    for name, size, before in (("Heading 1", 15, 14), ("Heading 2", 12.5, 10)):
        st = doc.styles[name]
        st.font.name = BODY_FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True


def build_toc_page(doc, entries):
    p_t = doc.add_paragraph()
    r = p_t.add_run("Оглавление")
    r.bold = True; r.font.size = Pt(16); r.font.color.rgb = ACCENT; r.font.name = BODY_FONT
    p_t.paragraph_format.space_after = Pt(12)
    for level, text, page in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        if level == 2:
            p.paragraph_format.left_indent = Cm(0.5)
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(16.4), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = p.add_run(text)
        r.font.size = Pt(10); r.font.name = BODY_FONT; r.bold = (level == 1)
        r2 = p.add_run("\t" + str(page))
        r2.font.size = Pt(10); r2.font.name = BODY_FONT
    doc.add_page_break()


def parse_md():
    with open(MD_PATH, encoding="utf-8") as f:
        lines = f.read().splitlines()

    title = None
    meta_lines = []
    blocks = []
    i = 0
    in_header_block = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1; continue
        if stripped.startswith("# ") and title is None:
            title = stripped[2:].strip(); i += 1; continue
        if stripped == "---" and in_header_block:
            in_header_block = False; i += 1; continue
        if in_header_block:
            meta_lines.append(stripped); i += 1; continue
        if stripped.startswith("### "):
            blocks.append(("h3", stripped[4:].strip())); i += 1; continue
        if stripped.startswith("## "):
            blocks.append(("h2", stripped[3:].strip())); i += 1; continue
        if stripped.startswith("# "):
            blocks.append(("h2", stripped[2:].strip())); i += 1; continue
        if stripped == "---":
            i += 1; continue
        if stripped.startswith("|"):
            tbl_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tbl_lines.append(lines[i].strip()); i += 1
            rows = []
            for tl in tbl_lines:
                cells = [c.strip() for c in tl.strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    continue
                rows.append(cells)
            if rows:
                blocks.append(("table", rows))
            continue
        m_ol = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m_ol:
            blocks.append(("oli", (m_ol.group(1), m_ol.group(2)))); i += 1; continue
        if stripped.startswith("- "):
            blocks.append(("li", stripped[2:].strip())); i += 1; continue
        blocks.append(("p", stripped)); i += 1
    return title, meta_lines, blocks


def build(out_path, toc_entries=None):
    title, meta_lines, blocks = parse_md()

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0); section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0); section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2); section.right_margin = Cm(2.2)
    style_setup(doc)
    add_footer_page_number(doc)

    cp = doc.core_properties
    cp.title = "Роналду vs Месси: глубокое сравнительное исследование"
    cp.author = "Z.ai"
    cp.subject = "Сравнительный анализ карьеры по официальной статистике (на 09.10.2026)"
    cp.creator = "Z.ai"

    build_title_page(doc, title, meta_lines)
    if toc_entries:
        build_toc_page(doc, toc_entries)

    # keep the final numbered list (appendix) together: chain all but its last item
    runs, cur = [], []
    for i, (k, _) in enumerate(blocks):
        if k == "oli" and (not cur or i == cur[-1] + 1):
            cur.append(i)
        else:
            if len(cur) > 1:
                runs.append(cur)
            cur = [i] if k == "oli" else []
    if len(cur) > 1:
        runs.append(cur)
    keep_oli = set(runs[-1][:-1]) if runs else set()

    for bidx, (kind, payload) in enumerate(blocks):
        if kind == "h2":
            doc.add_heading(payload, level=1)
        elif kind == "h3":
            doc.add_heading(payload, level=2)
        elif kind == "p":
            p = doc.add_paragraph()
            emit_inline(p, payload, size=Pt(11))
            # bold-only short paragraph = lead-in ("ЧМ-2026 (дополнительно):")
            # -> keep it on the same page as its list
            if re.fullmatch(r"\*\*[^*]+\*\*:?", payload) and len(payload) < 80:
                p.paragraph_format.keep_with_next = True
        elif kind == "li":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run("\u2022  "); r.font.size = Pt(11); r.font.name = BODY_FONT
            emit_inline(p, payload, size=Pt(11))
        elif kind == "oli":
            num, text = payload
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(2)
            if bidx in keep_oli:
                p.paragraph_format.keep_with_next = True
            r = p.add_run(f"{num}. "); r.bold = True
            r.font.size = Pt(11); r.font.name = BODY_FONT; r.font.color.rgb = ACCENT
            emit_inline(p, text, size=Pt(11))
        elif kind == "table":
            rows = payload
            ncols = max(len(rw) for rw in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table_full_width(table)
            table.autofit = True
            for ridx, row_cells in enumerate(rows):
                row = table.add_row()
                no_row_split(row)
                if ridx == 0:
                    mark_header_row(row)
                for cidx in range(ncols):
                    cell = row.cells[cidx]
                    text = row_cells[cidx] if cidx < len(row_cells) else ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.space_before = Pt(1)
                    if ridx == 0:
                        set_cell_shading(cell, "1F3A5F")
                        emit_inline(p, text, size=Pt(9), base_bold=True,
                                    color=RGBColor(0xFF, 0xFF, 0xFF))
                    else:
                        if ridx % 2 == 0:
                            set_cell_shading(cell, "EEF2F7")
                        emit_inline(p, text, size=Pt(9))
            # spacing paragraph after table
            sp = doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4)
            r = sp.add_run(); r.font.size = Pt(4)

    doc.save(out_path)
    print("saved:", out_path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else OUT_PATH
    toc = None
    if len(sys.argv) > 2:
        with open(sys.argv[2], encoding="utf-8") as f:
            toc = [tuple(e) for e in json.load(f)]
    build(out, toc)
